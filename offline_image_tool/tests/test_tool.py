import copy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import image_tool as tool

BASE = {"text": [], "tables": [], "flowcharts": [], "charts": [], "uncertainties": []}

class Tests(unittest.TestCase):
    def test_hash_matches_analyzed_bytes_when_source_changes(self):
        import hashlib
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp) / "image.png", Path(tmp) / "out"
            Image.new("RGB", (64, 64), "white").save(source)
            expected = hashlib.sha256(source.read_bytes()).hexdigest()
            class Fake:
                def generate(self, *args):
                    source.write_bytes(b"source changed during inference")
                    return json.dumps(BASE), False
            tool.analyze(Fake(), source, out)
            result = json.loads((out / "result.json").read_text())
            self.assertEqual(result["source_sha256"], expected)

    def test_flat_headings_normalized_without_changing_text(self):
        value = copy.deepcopy(BASE)
        value["tables"] = [{"title": "", "header_rows": ["Item", "Price"], "rows": [["Apples", "3.50"]]}]
        parsed = tool.parse_result(json.dumps(value))
        self.assertEqual(parsed["tables"][0]["header_rows"], [["Item", "Price"]])
        self.assertEqual(parsed["tables"][0]["rows"], [["Apples", "3.50"]])

    def test_missing_headers_field_rejected(self):
        value = copy.deepcopy(BASE)
        value["tables"] = [{"title": "bad", "rows": [["12"]]}]
        with self.assertRaises(ValueError):
            tool.parse_result(json.dumps(value))

    def test_focused_mode_allows_omitted_unrelated_categories(self):
        value = {"flowcharts": [{"title": "", "nodes": [{"id": "1", "label": "Start"}], "edges": []}]}
        parsed = tool.parse_result(json.dumps(value), "flowchart")
        self.assertEqual(parsed["tables"], [])
        self.assertIn("omitted", parsed["uncertainties"][0])
        with self.assertRaises(ValueError):
            tool.parse_result(json.dumps(value), "auto")

    def test_chart_focus_excludes_other_content_and_flags_axis(self):
        value = copy.deepcopy(BASE)
        value["tables"] = [{"invented": "unrequested data"}]
        value["charts"] = [{"title": "", "x_axis": {"label": "Month", "scale": "linear", "ticks": ["Jan", "Feb"]}, "y_axis": {"label": "Units", "scale": "linear", "ticks": ["0", "10"]}, "points": []}]
        parsed = tool.parse_result(json.dumps(value), "chart")
        self.assertEqual(parsed["tables"], [])
        self.assertTrue(any("misclassified" in warning for warning in parsed["uncertainties"]))
        self.assertEqual(parsed["charts"][0]["x_axis"]["scale"], "linear")

    def test_reject_ragged_table(self):
        value = copy.deepcopy(BASE)
        value["tables"] = [{"title": "bad", "header_rows": [["A", "B"]], "rows": [["1"]]}]
        with self.assertRaises(ValueError):
            tool.parse_result(json.dumps(value))

    def test_reject_dangling_arrow(self):
        value = copy.deepcopy(BASE)
        value["flowcharts"] = [{"title": "bad", "nodes": [{"id": "A", "label": "Start"}], "edges": [{"source": "A", "target": "B", "label": ""}]}]
        with self.assertRaises(ValueError):
            tool.parse_result(json.dumps(value))

    def test_reject_duplicate_node(self):
        value = copy.deepcopy(BASE)
        value["flowcharts"] = [{"title": "bad", "nodes": [{"id": "A", "label": "Start"}, {"id": "A", "label": "Stop"}], "edges": []}]
        with self.assertRaises(ValueError):
            tool.parse_result(json.dumps(value))

    def test_literal_excel_and_null(self):
        from openpyxl import load_workbook
        value = copy.deepcopy(BASE)
        value["tables"] = [{"title": "Test", "header_rows": [["Code", "Text", "Missing"]], "rows": [["0012", "=1+1", None]]}]
        result = {"data": value, "source_sha256": "test", "verification": "not requested", "warnings": []}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            tool.save_workbook(result, path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Table 1"])
            self.assertEqual(wb.active.title, "Table 1")
            self.assertEqual(wb["Table 1"]["A2"].value, "0012")
            self.assertEqual(wb["Table 1"]["B2"].value, "=1+1")
            self.assertEqual(wb["Table 1"]["B2"].data_type, "s")
            self.assertIsNone(wb["Table 1"]["C2"].value)
            wb.close()

    def test_color_table_round_trip_and_excel_fills(self):
        from openpyxl import load_workbook
        value = copy.deepcopy(BASE)
        value["tables"] = [{"title": "", "header_rows": [["Type", "Value"]],
                            "rows": [["π", "42"]],
                            "cell_colors": [["#FFF2CC", "#DDEBF7"], [None, "#E2F0D9"]]}]
        parsed = tool.parse_result(json.dumps(value), "table")
        self.assertEqual(parsed["tables"][0]["cell_colors"][0][0], "#FFF2CC")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "colors.xlsx"
            result = {"data": parsed}
            tool.save_workbook(result, path)
            wb = load_workbook(path)
            self.assertEqual(wb["Table 1"]["A1"].fill.fgColor.rgb[-6:], "FFF2CC")
            self.assertEqual(wb["Table 1"]["B2"].fill.fgColor.rgb[-6:], "E2F0D9")
            wb.close()

    def test_tt_pi_cell_gets_review_comment(self):
        from openpyxl import load_workbook
        value = copy.deepcopy(BASE)
        value["tables"] = [{"title": "", "header_rows": [["Symbol"]], "rows": [["TT"]]}]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ambiguous.xlsx"
            tool.save_workbook({"data": value}, path)
            wb = load_workbook(path)
            self.assertIn("TT and π", wb["Table 1"]["A2"].comment.text)
            wb.close()

    def test_reject_misaligned_cell_colors(self):
        value = copy.deepcopy(BASE)
        value["tables"] = [{"title": "", "header_rows": [["A", "B"]], "rows": [["1", "2"]],
                            "cell_colors": [["#FFFFFF"]]}]
        with self.assertRaises(ValueError):
            tool.parse_result(json.dumps(value), "table")

    def test_flowchart_png_preview_from_node_boxes(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            source, output = Path(tmp) / "source.png", Path(tmp) / "preview.png"
            Image.new("RGB", (80, 60), "white").save(source)
            created = tool.save_flowchart_preview(source, {"nodes": [
                {"id": "N1", "bbox": [10, 10, 50, 40]}]}, output)
            self.assertTrue(created)
            self.assertTrue(output.is_file())

    def test_fill_sampler_reads_dominant_cell_color(self):
        from PIL import Image, ImageDraw
        sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
        import ocr_engine
        image = Image.new("RGB", (40, 30), "#FFF2CC")
        ImageDraw.Draw(image).text((14, 10), "42", fill="black")
        self.assertEqual(ocr_engine._cell_fill(image, 0, 0, 40, 30), "#FFF2CC")

    def test_possible_tt_pi_is_flagged_without_replacing_text(self):
        from PIL import Image
        import ocr_engine
        original_ocr = ocr_engine._ocr
        ocr_engine._ocr = lambda image: [{"text": "TT", "score": 0.99,
            "left": 10.0, "top": 12.0, "right": 30.0, "bottom": 25.0,
            "cx": 20.0, "cy": 18.5}]
        try:
            data = ocr_engine.extract(Image.new("RGB", (40, 30), "white"), "Focus on text")
        finally:
            ocr_engine._ocr = original_ocr
        self.assertEqual(data["text"], ["TT"])
        self.assertTrue(any("TT/π" in warning for warning in data["uncertainties"]))

    def test_truncation_preserves_raw_without_workbook(self):
        from PIL import Image
        class Fake:
            def generate(self, *args):
                return '{"text":', True
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp) / "image.png", Path(tmp) / "out"
            Image.new("RGB", (64, 64), "white").save(source)
            with self.assertRaises(ValueError):
                tool.analyze(Fake(), source, out)
            self.assertTrue((out / "raw_response.txt").exists())
            self.assertFalse((out / "result.xlsx").exists())

    def test_existing_output_not_overwritten(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "image.png"
            Image.new("RGB", (64, 64), "white").save(source)
            with self.assertRaises(FileExistsError):
                tool.analyze(None, source, tmp)

    def test_disagreement_flag(self):
        from PIL import Image
        first, second = copy.deepcopy(BASE), copy.deepcopy(BASE)
        first["text"], second["text"] = ["123"], ["128"]
        class Fake:
            def __init__(self):
                self.values = iter([first, second])
            def generate(self, *args):
                return json.dumps(next(self.values)), False
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp) / "image.png", Path(tmp) / "out"
            Image.new("RGB", (64, 64), "white").save(source)
            result = tool.analyze(Fake(), source, out, verify=True)
            self.assertIn("disagree", result["verification"])

    def test_mcp_handshake_discovery_and_error(self):
        requests = [
            {"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {"protocolVersion": "2025-03-26"}},
            {"jsonrpc": "2.0", "method": "notifications/initialized"},
            {"jsonrpc": "2.0", "id": 2, "method": "tools/list"},
            {"jsonrpc": "2.0", "id": 3, "method": "tools/call", "params": {"name": "analyze_image", "arguments": {}}},
        ]
        completed = subprocess.run([sys.executable, str(tool.ROOT / "image_tool.py"), "mcp"], input="\n".join(map(json.dumps, requests)) + "\n", text=True, capture_output=True, check=True)
        replies = [json.loads(line) for line in completed.stdout.splitlines()]
        self.assertEqual(len(replies), 3)
        self.assertEqual(replies[0]["result"]["protocolVersion"], "2025-03-26")
        self.assertEqual(replies[1]["result"]["tools"][0]["name"], "analyze_image")
        self.assertTrue(replies[2]["result"]["isError"])

if __name__ == "__main__":
    unittest.main()
